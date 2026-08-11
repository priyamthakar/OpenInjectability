"""Inspect additional Figshare articles from keyword search (Ackermann SI, micro-capillary, etc.)."""

from __future__ import annotations

import json
import ssl
import urllib.request
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

CTX = ssl.create_default_context()
UA = {"User-Agent": "OpenInjectability/1.0"}
OUT_DIR = Path(__file__).resolve().parent / "_tmp_figshare"
OUT_DIR.mkdir(exist_ok=True)
REPORT = Path(__file__).resolve().parent / "_repo_deep_pull_figshare_extra.json"


def get_json(url: str):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, context=CTX, timeout=90) as r:
        return json.load(r)


def download(url: str, dest: Path):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, context=CTX, timeout=120) as r:
        dest.write_bytes(r.read())
    return dest.stat().st_size


def docx_text(path: Path, limit: int = 4000) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    texts = []
    for t in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
        if t.text:
            texts.append(t.text)
        if t.tail:
            texts.append(t.tail)
    full = " ".join(texts)
    import re

    full = re.sub(r"\s+", " ", full)
    return full[:limit]


def main():
    ids = [
        32195520,  # Felfeli raw force (already known)
        20800108,  # Rini
        31976504,  # Ackermann AF1
        31976507,  # Ackermann AF2
        31976510,  # Ackermann AF3
        32195514,  # Felfeli AF8
        32195496,  # Felfeli AF2
        27412368,  # Flow and injection characteristics micro-capillary
        32923061,  # Guiding syringe selection collection/article?
    ]
    report = []
    for aid in ids:
        entry = {"article_id": aid}
        try:
            data = get_json(f"https://api.figshare.com/v2/articles/{aid}")
            entry.update(
                {
                    "title": data.get("title"),
                    "doi": data.get("doi"),
                    "url": data.get("url_public_html"),
                    "defined_type": data.get("defined_type_name"),
                    "license": data.get("license"),
                    "description": (data.get("description") or "")[:800],
                    "files": [
                        {
                            "name": f.get("name"),
                            "size": f.get("size"),
                            "download_url": f.get("download_url"),
                            "mimetype": f.get("mimetype"),
                        }
                        for f in (data.get("files") or [])
                    ],
                }
            )
            print(f"\n=== {aid}: {entry['title'][:100]} ===")
            print("doi", entry["doi"], "nfiles", len(entry["files"]))
            for f in entry["files"]:
                print(" file", f["name"], f["size"], f["mimetype"])
                # download small files only for inspect
                name = f["name"] or "file.bin"
                dest = OUT_DIR / f"{aid}_{name}"
                if f.get("size") and f["size"] < 5_000_000 and f.get("download_url"):
                    try:
                        sz = download(f["download_url"], dest)
                        entry.setdefault("downloaded", []).append(
                            {"path": str(dest), "size": sz, "name": name}
                        )
                        low = name.lower()
                        if low.endswith(".docx"):
                            txt = docx_text(dest, 3500)
                            entry.setdefault("text_snips", {})[name] = txt
                            print("  docx snip:", txt[:400].replace("\n", " "))
                        elif low.endswith((".xlsx", ".xls", ".csv")):
                            if low.endswith(".csv"):
                                text = dest.read_text(encoding="utf-8", errors="replace")[:1500]
                                entry.setdefault("text_snips", {})[name] = text
                                print("  csv head:", text[:300].replace("\n", " | "))
                            else:
                                import openpyxl

                                wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
                                sheets = {}
                                for sn in wb.sheetnames[:3]:
                                    ws = wb[sn]
                                    rows = []
                                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                                        rows.append(
                                            [str(c) if c is not None else None for c in row[:15]]
                                        )
                                        if i >= 4:
                                            break
                                    sheets[sn] = rows
                                wb.close()
                                entry.setdefault("xlsx_preview", {})[name] = sheets
                                print("  xlsx sheets", list(sheets.keys()))
                                for sn, rows in sheets.items():
                                    print("   ", sn, "row0", rows[0] if rows else None)
                        elif low.endswith(".pdf"):
                            entry.setdefault("text_snips", {})[name] = (
                                f"PDF size={sz} (not text-extracted)"
                            )
                    # Per-file isolation is intentional for heterogeneous repository data.
                    except Exception as de:  # noqa: BLE001
                        entry.setdefault("download_errors", []).append(f"{name}: {de}")
                        print("  download err", de)
        # Per-record isolation keeps one remote/data-format failure from ending the batch.
        except Exception as e:  # noqa: BLE001
            entry["error"] = f"{type(e).__name__}: {e}"
            print("ERR", aid, e)
        report.append(entry)

    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print("\nWrote", REPORT)


if __name__ == "__main__":
    main()
